"""Local FPI sector-flow dashboard. Run: python3 app.py"""
from __future__ import annotations

import cgi
import hashlib
import html
import io
import json
import math
import os
import re
import sqlite3
import uuid
from datetime import date, datetime
from html.parser import HTMLParser
from pathlib import Path
from wsgiref.simple_server import make_server

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

BASE = Path(__file__).parent
DATA = BASE / "data"
RAW = DATA / "raw"
DB = DATA / "market.db"

from cloud.fpi.core import TableReader, parse_report


def conn():
    DATA.mkdir(exist_ok=True); RAW.mkdir(exist_ok=True)
    db=sqlite3.connect(DB); db.row_factory=sqlite3.Row
    db.executescript("""create table if not exists reports(id text primary key, report_date text unique, period_start text, period_end text, filename text, raw_path text, hash text unique, error text, imported_at text);
    create table if not exists flows(report_id text, sector text, net real, auc real, primary key(report_id,sector));""")
    return db


def import_report(db, raw, filename):
    digest=hashlib.sha256(raw).hexdigest()
    raw_path=RAW/f'{digest}.html'
    raw_path.write_bytes(raw)  # Retain even rejected uploads.
    report_date,start,end,flows=parse_report(raw)
    if db.execute('select 1 from reports where report_date=? or hash=?',(str(report_date),digest)).fetchone():
        raise ValueError('This fortnight is already in your dashboard')
    report_id=uuid.uuid4().hex
    with db:
        db.execute('insert into reports values(?,?,?,?,?,?,?,?,?)',(report_id,str(report_date),str(start),str(end),filename,str(raw_path),digest,None,datetime.now().isoformat()))
        db.executemany('insert into flows values(?,?,?,?)',[(report_id,*f) for f in flows])
    return {'id':report_id,'reportDate':str(report_date),'sectorCount':len(flows)}


def dashboard(db):
    reports=db.execute("select * from reports where error is null order by report_date desc").fetchall()
    if not reports: return "<section class='empty'><h2>Upload your first NSDL report</h2><p>Choose one fortnightly NSDL HTML file. The app stores it, validates it, and then shows the sector flow.</p></section>"
    latest=reports[0]; flows=db.execute("select * from flows where report_id=? order by net desc",(latest["id"],)).fetchall()
    total=sum(x["net"] for x in flows); pos=[x for x in flows if x["net"]>=0]; neg=sorted((x for x in flows if x["net"]<0),key=lambda x:x["net"])
    fmt=lambda x:f"{x:+,.0f}"
    rows="".join(f"<div class='row'><b>{html.escape(x['sector'])}</b><strong class={'up' if x['net']>=0 else 'down'}>{fmt(x['net'])} Cr</strong><small>Ending AUC {x['auc']:,.0f} Cr</small></div>" for x in flows)
    rank=lambda items:"".join(f"<div class='row'><span>{html.escape(x['sector'])}</span><strong class={'up' if x['net']>=0 else 'down'}>{fmt(x['net'])} Cr</strong></div>" for x in items[:5])
    return f"""<p class='period'>{latest['period_start']} to {latest['period_end']} · FPI · ₹ crore</p><section class='cards'><article><small>FPI net investment</small><strong class={'up' if total>=0 else 'down'}>{fmt(total)} Cr</strong></article><article><small>Money moving in</small><strong>{len(pos)} sectors</strong></article><article><small>Money moving out</small><strong>{len(neg)} sectors</strong></article></section><section class='ranks'><article><h2>Money moving in</h2>{rank(pos)}</article><article><h2>Money moving out</h2>{rank(neg)}</article></section><section class='all'><h2>All sector flows</h2><p>Net Investment is the fortnight’s FPI flow. AUC is FPI holdings at period end.</p>{rows}</section>"""


STYLE="""<style>*{box-sizing:border-box}body{margin:0;background:#f4f7f4;color:#13231e;font:16px system-ui,sans-serif}main{max-width:960px;margin:auto;padding:26px 16px 70px}header{display:flex;justify-content:space-between;gap:20px;align-items:start}h1{font-size:clamp(30px,6vw,50px);margin:6px 0;letter-spacing:-.05em}h2{font-size:19px}.eyebrow,.period,small,p{color:#64736b}.eyebrow{font-size:11px;font-weight:700;letter-spacing:.12em}.upload{background:#146c49;color:white;padding:12px 16px;border-radius:9px;font-weight:700;cursor:pointer}.upload input{display:none}.cards,.ranks{display:grid;gap:12px;margin:22px 0}.cards{grid-template-columns:repeat(3,1fr)}article,.all,.empty{background:#fff;border:1px solid #e2e9e3;border-radius:12px;padding:18px}.cards strong{display:block;font-size:25px;margin-top:7px}.ranks{grid-template-columns:1fr 1fr}.row{display:flex;gap:10px;align-items:center;justify-content:space-between;border-top:1px solid #edf1ed;padding:11px 0}.row small{margin-left:auto}.up{color:#07834d}.down{color:#c23d36}.error{padding:12px;background:#fff0ee;color:#9a302a;border-radius:8px}@media(max-width:650px){header{display:block}.upload{display:inline-block;margin-top:14px}.cards{grid-template-columns:1fr 1fr}.ranks{grid-template-columns:1fr}.row{flex-wrap:wrap}.row small{width:100%;margin-left:0}}</style>"""


def app(environ, start_response):
    db=conn(); path=environ["PATH_INFO"]
    try:
		# React dashboard API
        if path == "/api/reports" and environ["REQUEST_METHOD"] == "GET":
            rows=[dict(x) for x in db.execute("select * from reports where error is null order by report_date desc")]
            for r in rows:
                p=TableReader(); p.feed(Path(r['raw_path']).read_text())
                t=next(t for t in p.tables if len(t)>4 and any('Sectors' in row for row in t))
                col=[i for i in range(len(t[0])) if 'Net Investment' in t[0][i] and 'IN INR' in t[1][i] and t[2][i]=='Equity'][-1]
                total=next(row for row in t if 'Grand Total' in row)
                r['totalNet']=float(total[col].replace(',',''))
            start_response("200 OK", [("Content-Type","application/json"),("Access-Control-Allow-Origin","*")]); return [json.dumps(rows).encode()]
        if path in ("/api/flows", "/api/flows/latest") and environ["REQUEST_METHOD"] == "GET":
            query="""select f.*,r.report_date reportDate,r.period_start periodStart,r.period_end periodEnd from flows f join reports r on r.id=f.report_id order by r.report_date desc, f.net desc"""
            rows=[]
            for x in db.execute(query): rows.append({"id":x["report_id"]+"-"+x["sector"],"sectorName":x["sector"],"equityNetInvestmentCr":x["net"],"equityAucCr":x["auc"],"reportDate":x["reportDate"],"periodStart":x["periodStart"],"periodEnd":x["periodEnd"]})
            if path.endswith("latest") and rows:
                latest=rows[0]["reportDate"]; rows=[x for x in rows if x["reportDate"]==latest]
            start_response("200 OK", [("Content-Type","application/json"),("Access-Control-Allow-Origin","*")]); return [json.dumps(rows).encode()]
        if path == "/api/reports/import" and environ["REQUEST_METHOD"] == "POST":
            if int(environ.get('CONTENT_LENGTH') or 0) > 10*1024*1024: raise ValueError('Please choose a report smaller than 10 MB')
            form=cgi.FieldStorage(fp=environ["wsgi.input"],environ=environ,keep_blank_values=True); field=form["file"] if "file" in form else None
            raw=field.file.read() if field is not None and getattr(field,"file",None) else b""; filename=Path(getattr(field,"filename","") or "report.html").name
            payload=import_report(db,raw,filename)
            start_response("201 Created", [("Content-Type","application/json"),("Access-Control-Allow-Origin","*")]); return [json.dumps(payload).encode()]
        if path=="/" and environ["REQUEST_METHOD"]=="GET":
            message=environ.get("QUERY_STRING","").replace("message="," ").replace("+"," ")
            body=f"<main><header><div><div class='eyebrow'>INDIAN EQUITY · FPI FLOWS</div><h1>Where is FPI money moving?</h1><p>Simple fortnightly sector data from NSDL.</p></div><form method='post' action='/upload' enctype='multipart/form-data'><label class='upload'>Upload NSDL HTML<input name='file' type='file' accept='.html,text/html' onchange='this.form.submit()'></label></form></header>{('<p class=error>'+html.escape(message)+'</p>') if message else ''}{dashboard(db)}</main>"
            start_response("200 OK", [("Content-Type","text/html")]); return [("<!doctype html>"+STYLE+body).encode()]
        if path=="/upload" and environ["REQUEST_METHOD"]=="POST":
            form=cgi.FieldStorage(fp=environ["wsgi.input"],environ=environ,keep_blank_values=True)
            item=form.getfirst("file") if not isinstance(form.getfirst("file"), list) else None
            field=form["file"] if "file" in form else None
            raw=field.file.read() if field is not None and getattr(field,"file",None) else b""
            filename=Path(getattr(field,"filename","") or "report.html").name
            report_date,start,end,flows=parse_report(raw); digest=hashlib.sha256(raw).hexdigest(); existing=db.execute("select 1 from reports where report_date=? or hash=?",(str(report_date),digest)).fetchone()
            if existing: raise ValueError("This report has already been imported")
            report_id=uuid.uuid4().hex; raw_path=RAW/f"{report_date}-{digest[:8]}.html"; raw_path.write_bytes(raw)
            db.execute("insert into reports values(?,?,?,?,?,?,?,?,?)",(report_id,str(report_date),str(start),str(end),filename,str(raw_path),digest,None,datetime.now().isoformat()))
            db.executemany("insert into flows values(?,?,?,?)",[(report_id,*f) for f in flows]); db.commit()
            start_response("303 See Other",[("Location","/")]); return [b""]
        if path.startswith("/export/"):
            report_id=path.split("/")[-1]; r=db.execute("select * from reports where id=?",(report_id,)).fetchone(); flows=db.execute("select * from flows where report_id=?",(report_id,)).fetchall()
            if not r: raise ValueError("Report not found")
            wb=Workbook(); ws=wb.active; ws.title="Normalized FPI Flows"; ws.append(["Report Date","Period Start","Period End","Institution","Sector","Equity Net Investment (Cr)","Equity AUC (Cr)"])
            for cell in ws[1]: cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="146C49")
            for f in flows: ws.append([r['report_date'],r['period_start'],r['period_end'],"FPI",f['sector'],f['net'],f['auc']])
            ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
            for col,width in [('A',16),('B',16),('C',16),('D',14),('E',44),('F',30),('G',25)]: ws.column_dimensions[col].width=width
            for row in ws.iter_rows(min_row=2):
                for cell in row: cell.font=Font(name='Arial',size=11)
                row[5].number_format=row[6].number_format='#,##0;[Red](#,##0);0'
            out=io.BytesIO();wb.save(out);start_response("200 OK",[("Content-Type","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),("Content-Disposition",f"attachment; filename=fpi-flows-{r['report_date']}.xlsx")]);return [out.getvalue()]
        start_response("404 Not Found",[]); return [b"Not found"]
    except Exception as e:
        if path.startswith('/api/'):
            start_response('400 Bad Request', [('Content-Type','application/json'),('Access-Control-Allow-Origin','*')]); return [json.dumps({'error':str(e)}).encode()]
        start_response("303 See Other",[("Location","/?message="+str(e).replace(" ","+"))]); return [b""]
    finally: db.close()


if __name__ == "__main__":
    print("Open http://localhost:8000")
    make_server("127.0.0.1",8000,app).serve_forever()

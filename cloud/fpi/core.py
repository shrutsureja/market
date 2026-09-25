"""Portable parsing and XLSX generation shared by local WSGI and Workers."""
import io
import math
import re
from datetime import date, datetime
from html.parser import HTMLParser
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

PARSER_VERSION = "1"

MONTHS = {m: i for i, m in enumerate(("January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"), 1)}


class TableReader(HTMLParser):
    def __init__(self):
        super().__init__(); self.tables=[]; self.table=None; self.row=None; self.cell=None; self.text=[]; self.span=1
    def handle_starttag(self, tag, attrs):
        attrs=dict(attrs)
        if tag == "table": self.table=[]
        elif tag == "tr" and self.table is not None: self.row=[]
        elif tag in ("td", "th") and self.row is not None:
            self.cell=True; self.text=[]; self.span=int(attrs.get("colspan", "1") or 1)
    def handle_data(self, data):
        if self.cell: self.text.append(data)
    def handle_endtag(self, tag):
        if tag in ("td", "th") and self.cell:
            value=" ".join("".join(self.text).split())
            self.row.extend([value] * self.span); self.cell=False
        elif tag == "tr" and self.row is not None: self.table.append(self.row); self.row=None
        elif tag == "table" and self.table is not None: self.tables.append(self.table); self.table=None


def parse_date(value: str) -> date:
    value=value.strip().replace(",", "")
    for fmt in ("%B %d %Y", "%d %B %Y", "%d-%b-%Y", "%d/%m/%Y"):
        try: return datetime.strptime(value, fmt).date()
        except ValueError: pass
    raise ValueError(f"unrecognised date: {value}")


def parse_document(raw: bytes):
    p=TableReader(); p.feed(raw.decode("utf-8", "ignore"))
    table=next((t for t in p.tables if any('sectors' in [c.lower() for c in row] for row in t) and 'net investment' in ' '.join(' '.join(row) for row in t).lower()), None)
    if not table: raise ValueError("Expected NSDL sector table was not found")
    header_end=next(i for i,row in enumerate(table) if 'sectors' in [c.lower() for c in row])
    header=table[:header_end+1]
    sector_col=next((i for i,v in enumerate(header[-1]) if v.strip().lower()=="sectors"), None)
    net_cols=[]; auc_cols=[]
    for i in range(max(map(len,header))):
        cells=[r[i].lower() for r in header if i<len(r)]
        labels=' '.join(cells)
        direct_equity=cells.count('equity') >= 2 and 'mutual funds' not in labels
        if "net investment" in labels and "in inr cr" in labels and direct_equity: net_cols.append(i)
        if "auc as on" in labels and "in inr cr" in labels and direct_equity: auc_cols.append(i)
    if sector_col is None or not net_cols or not auc_cols: raise ValueError("Net Investment / Equity or AUC / Equity column not found")
    groups=[x for x in header[0] if "auc as on" in x.lower()]
    periods=[x for x in header[0] if "net investment" in x.lower()]
    if not groups or not periods: raise ValueError("Current report date or period was not found")
    report_date=parse_date(re.search(r"([A-Za-z]+\s+\d{1,2},?\s*\d{4})",groups[-1]).group(1))
    m=re.search(r"([A-Za-z]+)\s+(\d{1,2})\s*-\s*(\d{1,2}),?\s*(\d{4})", periods[-1])
    if not m: raise ValueError("Current Net Investment period was not recognised")
    start=date(int(m.group(4)), MONTHS[m.group(1)], int(m.group(2))); end=date(int(m.group(4)), MONTHS[m.group(1)], int(m.group(3)))
    if end != report_date: raise ValueError("Net Investment period and AUC report date do not reconcile")
    def number(x):
        x=x.replace(',', '').replace('₹', '').strip()
        if not x: raise ValueError('A required financial value is blank')
        value=float(x)
        if not math.isfinite(value): raise ValueError('Invalid financial value')
        return value
    flows=[]; total=None; names=set()
    for row in table[header_end+1:]:
        if not any(row): continue
        if len(row)<=max(sector_col,net_cols[-1],auc_cols[-1]): raise ValueError('Incomplete sector row; import stopped')
        sector=row[sector_col].strip()
        net,auc=number(row[net_cols[-1]]),number(row[auc_cols[-1]])
        if sector.lower()=='grand total': total=(net,auc); continue
        if not sector or sector.lower() in names: raise ValueError('Blank or duplicate sector name')
        names.add(sector.lower()); flows.append((sector,net,auc))
    if not 10 <= len(flows) <= 100: raise ValueError(f'Unexpected sector count: {len(flows)}')
    if total is None: raise ValueError('Grand Total not found')
    # Each source row is rounded to a whole crore: allow at most half a crore per row plus total rounding.
    for index, expected in enumerate(total,1):
        if abs(sum(f[index] for f in flows)-expected) > (len(flows)+1)*0.5:
            raise ValueError('Grand Total validation failed; import stopped')
    return dict(report_date=str(report_date), period_start=str(start), period_end=str(end), flows=flows, totalNet=total[0], totalAuc=total[1])


def parse_report(raw):
    doc = parse_document(raw)
    return (date.fromisoformat(doc['report_date']), date.fromisoformat(doc['period_start']),
            date.fromisoformat(doc['period_end']), doc['flows'])


def export_workbook(report, flows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Normalized FPI Flows"
    ws.append(["Report Date", "Period Start", "Period End", "Institution", "Sector", "Equity Net Investment (Cr)", "Equity AUC (Cr)"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="146C49")
    for sector, net, auc in flows:
        ws.append([report['report_date'], report['period_start'], report['period_end'], "FPI", sector, net, auc])
        # Imported sector names are always text, even if a source starts with '='.
        ws.cell(ws.max_row, 5).data_type = 's'
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for col, width in [('A',16),('B',16),('C',16),('D',14),('E',44),('F',30),('G',25)]:
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.font = Font(name='Arial', size=11)
        row[5].number_format = row[6].number_format = '#,##0;[Red](#,##0);0'
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
